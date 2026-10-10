"""Tests for the rag_ingestion document pipeline."""

from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.contenttypes.models import ContentType
from django.test import TestCase

from legal_kb.models import LegalKnowledgeEmbedding

from .interfaces import Chunk, EmbeddedChunk
from .models import DocumentSource
from .pipeline import EmptyDocumentError, IngestionPipeline, UnsupportedFormatError
from .registry import select_parser
from .services import IngestionService

User = get_user_model()

MARKDOWN = (
    '# SQL 注入防御\n\n'
    'SQL 注入通过拼接恶意 SQL 语句读取数据库数据。\n\n'
    '## 防御手段\n\n'
    '使用参数化查询（预编译语句）并对输入做白名单校验，可以有效防御 SQL 注入。\n'
).encode('utf-8')


class RegistryTests(TestCase):
    def test_plain_parser_selected_for_markdown(self):
        parser = select_parser('note.md', 'text/markdown')
        self.assertIsNotNone(parser)
        self.assertEqual(parser.name, 'plain')

    def test_html_is_stripped_to_text(self):
        parser = select_parser('page.html', 'text/html')
        parsed = parser.parse(b'<html><body><h1>Hi</h1><script>bad()</script><p>Body</p></body></html>',
                              filename='page.html', mime='text/html')
        self.assertIn('Hi', parsed.text)
        self.assertIn('Body', parsed.text)
        self.assertNotIn('bad()', parsed.text)


class IngestionServiceTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='kb-admin', password='pass12345', is_staff=True)

    def test_markdown_ingestion_writes_document_embeddings(self):
        document = IngestionService().ingest_bytes(
            data=MARKDOWN,
            filename='sql-defense.md',
            mime='text/markdown',
            uploaded_by=self.user,
        )

        self.assertEqual(document.status, 'completed')
        self.assertGreaterEqual(document.embedded_count, 1)
        content_type = ContentType.objects.get_for_model(DocumentSource)
        chunks = LegalKnowledgeEmbedding.objects.filter(
            source_type='document',
            content_type=content_type,
            object_id=document.id,
        )
        self.assertEqual(chunks.count(), document.embedded_count)
        self.assertTrue(chunks.filter(text__icontains='参数化查询').exists())
        self.assertEqual(chunks.first().metadata.get('visibility'), 'internal')

    def test_ingestion_is_idempotent_for_identical_content(self):
        service = IngestionService()
        first = service.ingest_bytes(data=MARKDOWN, filename='a.md', uploaded_by=self.user)
        second = service.ingest_bytes(data=MARKDOWN, filename='a.md', uploaded_by=self.user)

        self.assertEqual(first.id, second.id)
        content_type = ContentType.objects.get_for_model(DocumentSource)
        self.assertEqual(
            LegalKnowledgeEmbedding.objects.filter(
                source_type='document', content_type=content_type, object_id=first.id
            ).count(),
            first.embedded_count,
        )

    def test_empty_document_is_marked_failed(self):
        document = IngestionService().ingest_bytes(
            data=b'   \n  ', filename='empty.txt', uploaded_by=self.user,
        )
        self.assertEqual(document.status, 'failed')
        self.assertIn('EmptyDocumentError', document.error_message)

    def test_uploaded_documents_are_not_learner_facing_by_default(self):
        """'document' is outside the teaching-safe allowlist, so KnowledgeAgent
        and agent_runtime must not surface uploaded files to learners."""
        from legal_kb.services.knowledge_retrieval_service import SafeKnowledgeRetrievalService

        IngestionService().ingest_bytes(data=MARKDOWN, filename='sql.md', uploaded_by=self.user)
        items = SafeKnowledgeRetrievalService().retrieve(query='参数化查询 防御 SQL 注入', top_k=10)

        self.assertNotIn('document', {item['source_type'] for item in items})


class PipelineUnitTests(TestCase):
    def test_unsupported_format_raises(self):
        pipeline = IngestionPipeline(parser_selector=lambda filename, mime: None)
        with self.assertRaises(UnsupportedFormatError):
            pipeline.run(data=b'x', filename='mystery.xyz', document_id=1)

    def test_empty_text_raises(self):
        document = DocumentSource.objects.create(
            title='t', filename='t.txt', source_hash='hash-empty', status='running',
        )
        with self.assertRaises(EmptyDocumentError):
            IngestionPipeline().run(data=b'   ', filename='t.txt', document_id=document.id)


class RepositoryAtomicityTests(TestCase):
    """replace() must swap the index atomically: a mid-write failure keeps the old one."""

    def _embedded(self, order, text):
        chunk = Chunk(text=text, order=order, title='new', metadata={})
        return EmbeddedChunk(
            chunk=chunk, vector=[0.0] * 8, local_vector=[0.0] * 8,
            model='local-hash-v1', dimension=8,
        )

    def test_failed_rebuild_rolls_back_and_keeps_old_index(self):
        from .adapters.repository_legalkb import LegalKbChunkRepository

        doc = DocumentSource.objects.create(
            title='t', filename='t.md', source_hash='atomic-1', status='completed',
        )
        content_type = ContentType.objects.get_for_model(DocumentSource)
        LegalKnowledgeEmbedding.objects.create(
            source_type='document', content_type=content_type, object_id=doc.id,
            chunk_index=0, title='old', text='OLD_COMPLETE_INDEX', text_hash='oldhash',
            embedding=[0.0], embedding_model='local-hash-v1', embedding_dimension=1, metadata={},
        )

        embedded = [self._embedded(1, 'NEW_1'), self._embedded(2, 'NEW_2')]
        real_create = LegalKnowledgeEmbedding.objects.create
        state = {'n': 0}

        def failing_create(**kwargs):
            state['n'] += 1
            if state['n'] == 2:
                raise RuntimeError('simulated disk failure on 2nd chunk')
            return real_create(**kwargs)

        with patch.object(LegalKnowledgeEmbedding.objects, 'create', side_effect=failing_create):
            with self.assertRaises(RuntimeError):
                LegalKbChunkRepository().replace(
                    document_id=doc.id, source_type='document',
                    embedded=embedded, base_metadata={'title': 't'},
                )

        remaining = list(
            LegalKnowledgeEmbedding.objects
            .filter(content_type=content_type, object_id=doc.id)
            .values_list('text', flat=True)
        )
        self.assertEqual(remaining, ['OLD_COMPLETE_INDEX'])  # no partial NEW_1 left behind


class MarkItDownParserTests(TestCase):
    """markitdown adapter must close the temp file before convert reopens it."""

    def test_tempfile_closed_before_convert_and_cleaned_up(self):
        import os
        import sys
        import types

        from .adapters.parser_markitdown import MarkItDownParser

        captured = {}

        class FakeResult:
            def __init__(self, text):
                self.text_content = text

        class FakeMarkItDown:
            def convert(self, path):
                captured['path'] = path
                # Reopen by path — would raise PermissionError on Windows if the
                # writer handle were still open. Succeeds because parse() closes it.
                with open(path, 'rb') as handle:
                    data = handle.read()
                return FakeResult('converted:' + data.decode('utf-8', 'ignore'))

        fake_module = types.ModuleType('markitdown')
        fake_module.MarkItDown = FakeMarkItDown

        parser = MarkItDownParser()
        parser._available = True  # pretend the optional dependency is installed

        with patch.dict(sys.modules, {'markitdown': fake_module}):
            parsed = parser.parse(b'hello pptx', filename='deck.pptx', mime='')

        self.assertIn('converted:hello pptx', parsed.text)
        self.assertFalse(os.path.exists(captured['path']))  # temp file cleaned up


class DocxContentTests(TestCase):
    """Preserve mixed paragraphs and nested tables in both parser paths."""

    @staticmethod
    def _document_bytes():
        import io
        import zipfile

        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            archive.writestr('[Content_Types].xml',
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Override PartName="/word/document.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                '</Types>')
            archive.writestr('_rels/.rels',
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
                'Target="word/document.xml"/></Relationships>')
            archive.writestr('word/document.xml',
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:body><w:p><w:r><w:t>INTRODUCTION</w:t></w:r></w:p>'
                '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>TABLE_ONLY_EVIDENCE &amp; POLICY</w:t></w:r></w:p>'
                '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>NESTED_TABLE</w:t></w:r></w:p>'
                '</w:tc></w:tr></w:tbl></w:tc></w:tr></w:tbl>'
                '<w:p><w:r><w:t>CONCLUSION</w:t></w:r></w:p></w:body></w:document>')
        return stream.getvalue()

    def _assert_content(self, text):
        self.assertEqual(text.splitlines(), [
            'INTRODUCTION', 'TABLE_ONLY_EVIDENCE & POLICY', 'NESTED_TABLE', 'CONCLUSION',
        ])

    def test_mixed_paragraphs_and_nested_tables_preserve_order(self):
        from .adapters.parser_docx import DocxParser

        result = DocxParser().parse(self._document_bytes(), filename='mixed.docx', mime='')
        self._assert_content(result.text)

    def test_python_docx_path_reads_table_paragraphs(self):
        from .adapters.parser_docx import DocxParser

        try:
            from docx import Document  # noqa: F401
        except ImportError:
            self.skipTest('python-docx unavailable; XML fallback tested separately')
        text = DocxParser()._with_python_docx(self._document_bytes())
        self.assertIsNotNone(text)
        self._assert_content(text)

    def test_xml_fallback_reads_mixed_content_and_decodes_entities(self):
        from .adapters.parser_docx import DocxParser

        parser = DocxParser()
        with patch.object(parser, '_with_python_docx', return_value=None):
            result = parser.parse(self._document_bytes(), filename='mixed.docx', mime='')
        self._assert_content(result.text)

    def test_mixed_docx_table_evidence_is_indexed(self):
        from django.test import override_settings

        with override_settings(RAG_INGESTION_USE_LANGCHAIN_SPLITTER=False):
            document = IngestionService().ingest_bytes(
                data=self._document_bytes(), filename='mixed.docx',
            )
        self.assertEqual(document.status, 'completed')
        self.assertTrue(LegalKnowledgeEmbedding.objects.filter(
            source_type='document', object_id=document.id,
            text__contains='TABLE_ONLY_EVIDENCE & POLICY',
        ).exists())


class IngestionOwnershipTests(TestCase):
    """Deterministic interleavings fence old attempts at both index and state."""

    DATA = b'Identical file uploaded by overlapping ingestion tasks'

    def _newer_upload(self, *, reingest=True):
        from .interfaces import ParsedDocument

        class Parser:
            name = 'newer-parser'

            def parse(self, data, **kwargs):
                return ParsedDocument(text='NEWER_INDEX')

        pipeline = IngestionPipeline(parser_selector=lambda filename, mime: Parser())
        return IngestionService(pipeline=pipeline).ingest_bytes(
            data=self.DATA, filename='newer.txt', visibility='staff', reingest=reingest,
        )

    def _assert_newer_wins(self, document):
        document.refresh_from_db()
        self.assertEqual(DocumentSource.objects.count(), 1)
        self.assertEqual(document.status, 'completed')
        self.assertEqual(document.error_message, '')
        self.assertEqual(document.filename, 'newer.txt')
        self.assertEqual(document.visibility, 'staff')
        self.assertEqual(document.parser, 'newer-parser')
        self.assertEqual(document.chunk_count, 1)
        self.assertEqual(document.embedded_count, 1)
        texts = list(LegalKnowledgeEmbedding.objects.filter(
            source_type='document', object_id=document.id,
        ).values_list('text', flat=True))
        self.assertEqual(texts, ['NEWER_INDEX'])

    def test_older_failure_does_not_clobber_newer_first_upload(self):
        from unittest.mock import Mock

        def older_run(**kwargs):
            self.assertEqual(self._newer_upload(reingest=False).status, 'completed')
            raise RuntimeError('older task fails after newer task succeeds')

        pipeline = Mock()
        pipeline.run.side_effect = older_run
        service = IngestionService(pipeline=pipeline)
        with patch.object(service, '_audit') as audit:
            document = service.ingest_bytes(data=self.DATA, filename='older.txt')
        self._assert_newer_wins(document)
        audit.assert_not_called()

    def test_older_success_cannot_replace_newer_index(self):
        from .adapters.embedder_legalkb import LegalKbEmbedder
        from .interfaces import ParsedDocument

        class OlderParser:
            name = 'older-parser'

            def parse(self, data, **kwargs):
                return ParsedDocument(text='OLDER_INDEX')

        embedder = LegalKbEmbedder()
        real_embed = embedder.embed

        def older_embed(chunks):
            self.assertEqual(self._newer_upload().status, 'completed')
            return real_embed(chunks)

        pipeline = IngestionPipeline(
            parser_selector=lambda filename, mime: OlderParser(), embedder=embedder,
        )
        with patch.object(embedder, 'embed', side_effect=older_embed):
            document = IngestionService(pipeline=pipeline).ingest_bytes(
                data=self.DATA, filename='older.txt',
            )
        self._assert_newer_wins(document)

    def test_older_success_cannot_overwrite_newer_statistics(self):
        from dataclasses import replace
        from unittest.mock import Mock

        actual_pipeline = IngestionPipeline()

        def older_run(**kwargs):
            outcome = actual_pipeline.run(**kwargs)
            self.assertEqual(self._newer_upload().status, 'completed')
            return replace(outcome, parser='older-parser', chunk_count=99, embedded_count=99)

        pipeline = Mock()
        pipeline.run.side_effect = older_run
        service = IngestionService(pipeline=pipeline)
        with patch.object(service, '_audit') as audit:
            document = service.ingest_bytes(data=self.DATA, filename='older.txt')
        self._assert_newer_wins(document)
        audit.assert_not_called()

    def test_reingestion_uses_a_new_task_token(self):
        from django.test import override_settings

        with override_settings(RAG_INGESTION_USE_LANGCHAIN_SPLITTER=False):
            first = IngestionService().ingest_bytes(data=self.DATA, filename='a.txt')
            token = first.ingestion_token
            second = IngestionService().ingest_bytes(data=self.DATA, filename='a.txt')
            self.assertEqual(second.ingestion_token, token)
            third = IngestionService().ingest_bytes(
                data=self.DATA, filename='a.txt', reingest=True,
            )
        self.assertIsNotNone(token)
        self.assertNotEqual(third.ingestion_token, token)
        self.assertEqual(third.status, 'completed')


    def test_older_success_cannot_override_a_newer_failed_rebuild(self):
        from unittest.mock import Mock
        from .adapters.embedder_legalkb import LegalKbEmbedder
        from .interfaces import ParsedDocument

        original = IngestionService().ingest_bytes(data=self.DATA, filename='original.txt')
        original_texts = list(LegalKnowledgeEmbedding.objects.filter(
            source_type='document', object_id=original.id,
        ).values_list('text', flat=True))

        class OlderParser:
            name = 'older-parser'

            def parse(self, data, **kwargs):
                return ParsedDocument(text='STALE_SUCCESS_MUST_NOT_COMMIT')

        embedder = LegalKbEmbedder()
        real_embed = embedder.embed

        def interleaved_embed(chunks):
            broken = Mock()
            broken.run.side_effect = RuntimeError('newer parser failure')
            newer = IngestionService(pipeline=broken).ingest_bytes(
                data=self.DATA, filename='newer.txt', visibility='staff', reingest=True,
            )
            self.assertEqual(newer.status, 'failed')
            return real_embed(chunks)

        pipeline = IngestionPipeline(
            parser_selector=lambda filename, mime: OlderParser(), embedder=embedder,
        )
        with patch.object(embedder, 'embed', side_effect=interleaved_embed):
            document = IngestionService(pipeline=pipeline).ingest_bytes(
                data=self.DATA, filename='older.txt', reingest=True,
            )
        self.assertEqual(document.status, 'failed')
        self.assertEqual(document.filename, 'newer.txt')
        self.assertEqual(document.visibility, 'staff')
        self.assertIn('newer parser failure', document.error_message)
        texts = list(LegalKnowledgeEmbedding.objects.filter(
            source_type='document', object_id=document.id,
        ).values_list('text', flat=True))
        self.assertEqual(texts, original_texts)
