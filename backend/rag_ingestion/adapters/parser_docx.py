"""Word .docx parser preserving body and table paragraphs in document order."""

import io
import zipfile
from typing import Set
from xml.etree import ElementTree

from ..interfaces import ParsedDocument

_WORD_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


class DocxParser:
    """Extract body text, including nested table cells, with an XML fallback."""

    name = 'docx'
    EXTENSIONS: Set[str] = {'.docx'}

    def supports(self, *, filename: str, mime: str) -> bool:
        lowered = (filename or '').lower()
        if any(lowered.endswith(ext) for ext in self.EXTENSIONS):
            return True
        return 'wordprocessingml' in (mime or '')

    def parse(self, data: bytes, *, filename: str, mime: str) -> ParsedDocument:
        text = self._with_python_docx(data)
        if not text or not text.strip():
            text = self._with_zipfile(data)
        return ParsedDocument(text=text or '', metadata={'parser': self.name})

    def _with_python_docx(self, data: bytes):
        try:
            import docx
            from docx.oxml.ns import qn
            from docx.text.paragraph import Paragraph

            document = docx.Document(io.BytesIO(data))
            # document.paragraphs excludes tables. Walk all body paragraphs so
            # mixed content and nested table cells retain their original order.
            paragraphs = (
                Paragraph(element, document).text
                for element in document.element.body.iter(qn('w:p'))
            )
            return '\n'.join(text for text in paragraphs if text.strip())
        except Exception:
            return None

    def _with_zipfile(self, data: bytes) -> str:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                document = ElementTree.fromstring(archive.read('word/document.xml'))
            paragraphs = []
            for paragraph in document.iter(f'{{{_WORD_NS}}}p'):
                fragments = []
                for element in paragraph.iter():
                    if element.tag == f'{{{_WORD_NS}}}t':
                        fragments.append(element.text or '')
                    elif element.tag == f'{{{_WORD_NS}}}tab':
                        fragments.append('\t')
                    elif element.tag in (f'{{{_WORD_NS}}}br', f'{{{_WORD_NS}}}cr'):
                        fragments.append('\n')
                text = ''.join(fragments)
                if text.strip():
                    paragraphs.append(text)
            return '\n'.join(paragraphs)
        except (OSError, ValueError, KeyError, zipfile.BadZipFile, ElementTree.ParseError):
            return ''
