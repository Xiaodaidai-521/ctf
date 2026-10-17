"""Bounded, version-bound table parsing. This module performs no model calls."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import time
import unicodedata
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import PurePosixPath
from typing import BinaryIO, Callable, Protocol
from xml.etree import ElementTree as ET

PARSER_VERSION = 'internal-table-v2'


PARSING_ERROR_CODES = frozenset({
    'CONFIG_INVALID', 'CONFIG_UNSUPPORTED', 'ENCODING_UNSUPPORTED', 'ENCODING_INVALID',
    'SCOPE_INVALID', 'SCOPE_MERGE_ANCHOR_MISSING', 'FINGERPRINT_TYPE_UNSUPPORTED',
    'OBJECT_NOT_FOUND', 'ASSET_UNAVAILABLE', 'ASSET_VERSION_MISMATCH', 'RESOURCE_LIMIT',
    'TASK_CANCELLED', 'LOCATION_MISMATCH', 'PARSE_FAILED', 'EMPTY_ASSET',
    'HEADER_OUT_OF_RANGE', 'UNSAFE_ARCHIVE', 'UNSAFE_XML', 'UNSUPPORTED_FORMAT',
    'UNSUPPORTED_FORMULA', 'UNSUPPORTED_CELL_TYPE', 'CELL_ERROR', 'MERGED_REGIONS_OVERLAP',
})


class InputError(ValueError):
    def __init__(self, code: str):
        self.code = code
        super().__init__(code)  # No source text, filename or parser exception in errors.


@dataclass(frozen=True)
class Limits:
    max_bytes: int = 20 * 1024 * 1024
    max_rows: int = 100_000
    max_columns: int = 256
    max_cells: int = 200_000
    max_cell_codepoints: int = 32_768
    max_zip_entries: int = 2_000
    max_uncompressed_bytes: int = 100 * 1024 * 1024
    max_seconds: int = 300

    def __post_init__(self):
        if any(type(v) is not int or v <= 0 for v in asdict(self).values()):
            raise InputError('CONFIG_INVALID')


@dataclass(frozen=True)
class ParserConfig:
    encoding: str = 'utf-8-sig'
    delimiter: str = ','
    quotechar: str = '"'
    doublequote: bool = True
    header_row: int | None = 1
    include_header: bool = False
    hidden_policy: str = 'include'
    formula_policy: str = 'formula_text_only'
    normalization: str = 'identity-v1'

    def __post_init__(self):
        if self.encoding not in ('utf-8', 'utf-8-sig', 'gb18030'):
            raise InputError('ENCODING_UNSUPPORTED')
        if (len(self.delimiter) != 1 or len(self.quotechar) != 1
                or self.delimiter in '\r\n\0' or self.quotechar in '\r\n\0'
                or self.delimiter == self.quotechar):
            raise InputError('CONFIG_INVALID')
        if self.header_row is not None and (type(self.header_row) is not int or self.header_row < 1):
            raise InputError('CONFIG_INVALID')
        if type(self.include_header) is not bool or type(self.doublequote) is not bool:
            raise InputError('CONFIG_INVALID')
        if self.hidden_policy not in ('include', 'exclude'):
            raise InputError('CONFIG_INVALID')
        # Cache values are observed and reported but are not authoritative scan input.
        if self.formula_policy != 'formula_text_only' or self.normalization not in ('identity-v1', 'nfkc-mapped-v1'):
            raise InputError('CONFIG_UNSUPPORTED')


@dataclass(frozen=True)
class Selection:
    sheet_id: str
    rows: tuple[tuple[int, int], ...]
    columns: tuple[tuple[int, int], ...]


@dataclass(frozen=True)
class Scope:
    asset_version_id: str
    selections: tuple[Selection, ...]
    include_header: bool = False


def _ranges(ranges):
    ordered = []
    for pair in ranges:
        if len(pair) != 2 or any(type(v) is not int or v < 1 for v in pair) or pair[0] > pair[1]:
            raise InputError('SCOPE_INVALID')
        ordered.append(tuple(pair))
    if not ordered:
        raise InputError('SCOPE_INVALID')
    merged = []
    for start, end in sorted(ordered):
        if merged and start <= merged[-1][1] + 1:
            merged[-1] = (merged[-1][0], max(end, merged[-1][1]))
        else:
            merged.append((start, end))
    return tuple(merged)


def normalize_scope(scope: Scope) -> Scope:
    if not scope.asset_version_id or not scope.selections or type(scope.include_header) is not bool:
        raise InputError('SCOPE_INVALID')
    selected = []
    seen = set()
    for item in scope.selections:
        if item.sheet_id in seen or not re.fullmatch(r'(csv:0|xlsx:[0-9]+)', item.sheet_id):
            raise InputError('SCOPE_INVALID')
        seen.add(item.sheet_id)
        selected.append(Selection(item.sheet_id, _ranges(item.rows), _ranges(item.columns)))
    return Scope(scope.asset_version_id, tuple(sorted(selected, key=lambda s: s.sheet_id)), scope.include_header)


def fingerprint(value) -> str:
    """Internal integer/string JSON fingerprint; floats intentionally not supported."""
    def check(v):
        if v is None or type(v) in (str, bool, int):
            return
        if isinstance(v, (list, tuple)):
            for item in v:
                check(item)
            return
        if isinstance(v, dict) and all(type(k) is str for k in v):
            for item in v.values():
                check(item)
            return
        raise InputError('FINGERPRINT_TYPE_UNSUPPORTED')
    check(value)
    raw = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)
    return 'sha256:' + hashlib.sha256(raw.encode('utf-8')).hexdigest()


@dataclass(frozen=True)
class AssetRef:
    version_id: str
    domain_id: str
    content_hash: str
    format: str
    byte_size: int
    private_ref: str
    state: str = 'ready'


class PrivateSource(Protocol):
    """An object-owner port, not a client-provided path or URL.

    The implementation MUST reauthorize principal/domain/version and supply a
    stable read-only snapshot. No production implementation is wired here.
    """
    def open_version(self, ref: AssetRef, *, principal: str, domain_id: str) -> BinaryIO: ...


def read_private(source: PrivateSource, ref: AssetRef, *, principal: str,
                 domain_id: str, limits: Limits, cancelled: Callable[[], bool] = lambda: False) -> bytes:
    if not principal or domain_id != ref.domain_id:
        raise InputError('OBJECT_NOT_FOUND')
    if ref.state != 'ready':
        raise InputError('ASSET_UNAVAILABLE')
    if type(ref.byte_size) is not int or ref.byte_size < 0 or ref.byte_size > limits.max_bytes:
        raise InputError('RESOURCE_LIMIT')
    budget = Guard(limits, cancelled)
    try:
        with source.open_version(ref, principal=principal, domain_id=domain_id) as stream:
            chunks = []
            size = 0
            while True:
                budget.check()
                block = stream.read(min(64 * 1024, limits.max_bytes - size + 1))
                if not block:
                    break
                size += len(block)
                if size > limits.max_bytes:
                    raise InputError('RESOURCE_LIMIT')
                chunks.append(block)
    except InputError:
        raise
    except Exception:
        raise InputError('ASSET_UNAVAILABLE') from None
    data = b''.join(chunks)
    verify_asset(data, ref, limits)
    return data


def verify_asset(data, ref, limits):
    if len(data) > limits.max_bytes:
        raise InputError('RESOURCE_LIMIT')
    if ref.state != 'ready':
        raise InputError('ASSET_UNAVAILABLE')
    actual = 'sha256:' + hashlib.sha256(data).hexdigest()
    if len(data) != ref.byte_size or actual != ref.content_hash:
        raise InputError('ASSET_VERSION_MISMATCH')


class Guard:
    def __init__(self, limits, cancelled):
        self.limits = limits
        self.cancelled = cancelled
        self.started = time.monotonic()

    def check(self):
        if self.cancelled():
            raise InputError('TASK_CANCELLED')
        if time.monotonic() - self.started > self.limits.max_seconds:
            raise InputError('RESOURCE_LIMIT')

    def text(self, text):
        self.check()
        if len(text) > self.limits.max_cell_codepoints:
            raise InputError('RESOURCE_LIMIT')


@dataclass(frozen=True)
class Cell:
    sheet_id: str
    row: int
    column: int
    source_text: str = field(repr=False)
    value_type: str
    channel: str = 'value'
    cache_status: str = 'not_applicable'
    number_format: str | None = None
    normalization: str = 'identity-v1'

    def mapping(self):
        """Return text and proven boundaries only; never guess offsets after composition.

        Combining sequences are indivisible. Cross-sequence compositions (e.g.
        Hangul Jamo) fail the whole-string check and disable character operations.
        """
        if self.normalization == 'identity-v1':
            return self.source_text, tuple((i, i + 1) for i in range(len(self.source_text)))
        if self.normalization != 'nfkc-mapped-v1':
            raise InputError('CONFIG_UNSUPPORTED')
        groups = []
        for index, char in enumerate(self.source_text):
            if groups and unicodedata.combining(char):
                groups[-1][1] = index + 1
            else:
                groups.append([index, index + 1])
        pieces, spans = [], []
        for start, end in groups:
            value = unicodedata.normalize('NFKC', self.source_text[start:end])
            pieces.append(value)
            spans.extend([(start, end)] * len(value))
        normalized = unicodedata.normalize('NFKC', self.source_text)
        if ''.join(pieces) != normalized:
            return normalized, None
        return normalized, tuple(spans)

    @property
    def normalized_text(self):
        return self.mapping()[0]

    @property
    def location_reliable(self):
        return self.mapping()[1] is not None

    def source_span(self, start: int, end: int) -> tuple[int, int]:
        text, mapping = self.mapping()
        if (mapping is None or type(start) is not int or type(end) is not int
                or not 0 <= start < end <= len(text)):
            raise InputError('LOCATION_MISMATCH')
        if ((start > 0 and mapping[start - 1] == mapping[start])
                or (end < len(mapping) and mapping[end - 1] == mapping[end])):
            raise InputError('LOCATION_MISMATCH')  # Do not split a ligature or combining sequence.
        return mapping[start][0], mapping[end - 1][1]


@dataclass(frozen=True)
class ParseResult:
    asset_version_id: str
    input_content_hash: str
    parser_version: str
    config_json: str
    config_hash: str
    scope: Scope
    scope_hash: str
    cells: tuple[Cell, ...]
    structure_json: str
    coverage_json: str
    warnings: tuple[str, ...]
    excluded_contents: tuple[str, ...]

    def __repr__(self):
        return f'ParseResult(parser={self.parser_version!r}, cells={len(self.cells)})'


def _within(number, ranges):
    return any(start <= number <= end for start, end in ranges)


def _csv(data, config, guard):
    try:
        text = data.decode(config.encoding, errors='strict')
    except UnicodeError:
        raise InputError('ENCODING_INVALID') from None
    if '\0' in text:
        raise InputError('PARSE_FAILED')
    if text.startswith('\ufeff'):
        raise InputError('ENCODING_INVALID')  # Require explicit BOM-aware encoding.
    rows = []
    width = None
    try:
        reader = csv.reader(io.StringIO(text, newline=''), delimiter=config.delimiter,
                            quotechar=config.quotechar, doublequote=config.doublequote, strict=True)
        for row in reader:
            guard.check()
            if len(rows) >= guard.limits.max_rows or len(row) > guard.limits.max_columns:
                raise InputError('RESOURCE_LIMIT')
            if width is None:
                width = len(row)
            if not row or len(row) != width:
                raise InputError('PARSE_FAILED')
            if (len(rows) + 1) * width > guard.limits.max_cells:
                raise InputError('RESOURCE_LIMIT')
            for value in row:
                guard.text(value)
            rows.append(tuple(row))
    except csv.Error:
        raise InputError('PARSE_FAILED') from None
    if not rows:
        raise InputError('EMPTY_ASSET')
    return [('csv:0', rows, {})], []


_XML_NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
_COORD = re.compile(r'^([A-Z]+)([1-9][0-9]*)$')


def _coordinate(raw):
    m = _COORD.fullmatch(raw or '')
    if not m:
        raise InputError('PARSE_FAILED')
    column = 0
    for ch in m[1]:
        column = column * 26 + ord(ch) - 64
    return int(m[2]), column


def _xlsx(data, config, guard):
    """Guard ZIP/XML before openpyxl can allocate workbook structures.

    No extraction to disk, macro execution, formula evaluation or network reads.
    Formula caches are metadata observations, never assumed fresh.
    """
    members = {}
    raw_cells = {}
    rectangle_cells = 0
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > guard.limits.max_zip_entries:
                raise InputError('RESOURCE_LIMIT')
            if sum(e.file_size for e in entries) > guard.limits.max_uncompressed_bytes:
                raise InputError('RESOURCE_LIMIT')
            actual = 0
            xml_nodes = 0
            for entry in entries:
                guard.check()
                name = entry.filename
                if name in members or '\\' in name or ':' in name or name.startswith('/') or '..' in PurePosixPath(name).parts:
                    raise InputError('UNSAFE_ARCHIVE')
                if entry.flag_bits & 1 or 'vbaproject' in name.lower():
                    raise InputError('UNSUPPORTED_FORMAT')
                blocks = []
                with archive.open(entry) as stream:
                    while True:
                        guard.check()
                        block = stream.read(min(64 * 1024, guard.limits.max_uncompressed_bytes - actual + 1))
                        if not block:
                            break
                        actual += len(block)
                        if actual > guard.limits.max_uncompressed_bytes:
                            raise InputError('RESOURCE_LIMIT')
                        blocks.append(block)
                value = b''.join(blocks)
                members[name] = value
                if name.endswith(('.xml', '.rels')):
                    # Reject alternate encodings to make entity rejection unambiguous.
                    if b'\x00' in value or b'<!DOCTYPE' in value.upper() or b'<!ENTITY' in value.upper():
                        raise InputError('UNSAFE_XML')
                    root = ET.fromstring(value)
                    for element in root.iter():
                        xml_nodes += 1
                        if xml_nodes > guard.limits.max_cells * 16 + 10_000:
                            raise InputError('RESOURCE_LIMIT')
                        guard.check()
                        if element.text:
                            # Allow XML manifests larger than a cell; actual cell text is checked below.
                            if len(element.text) > guard.limits.max_uncompressed_bytes:
                                raise InputError('RESOURCE_LIMIT')
                    if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                        cells = {}
                        rows_seen = set()
                        height, width = 0, 0
                        for row_element in root.iter(_XML_NS + 'row'):
                            row_num = int(row_element.get('r', '0'))
                            if row_num in rows_seen or row_num < 1:
                                raise InputError('PARSE_FAILED')
                            rows_seen.add(row_num)
                            if row_num > guard.limits.max_rows:
                                raise InputError('RESOURCE_LIMIT')
                            for cell in row_element.findall(_XML_NS + 'c'):
                                position = _coordinate(cell.get('r'))
                                if position in cells or position[0] != row_num:
                                    raise InputError('PARSE_FAILED')
                                if position[1] > guard.limits.max_columns:
                                    raise InputError('RESOURCE_LIMIT')
                                height, width = max(height, position[0]), max(width, position[1])
                                formula = cell.find(_XML_NS + 'f')
                                cached = cell.find(_XML_NS + 'v')
                                cells[position] = (cell.get('t', 'n'), cached.text if cached is not None else None, formula is not None)
                                if len(cells) > guard.limits.max_cells:
                                    raise InputError('RESOURCE_LIMIT')
                        merged_positions = set()
                        for merged in root.iter(_XML_NS + 'mergeCell'):
                            endpoints = merged.get('ref', '').split(':')
                            start, end = endpoints if len(endpoints) == 2 else (endpoints[0], endpoints[0])
                            r1, c1 = _coordinate(start)
                            r2, c2 = _coordinate(end)
                            if r1 > r2 or c1 > c2:
                                raise InputError('PARSE_FAILED')
                            if r2 > guard.limits.max_rows or c2 > guard.limits.max_columns or (r2-r1+1)*(c2-c1+1) > guard.limits.max_cells:
                                raise InputError('RESOURCE_LIMIT')
                            height, width = max(height, r2), max(width, c2)
                            for mr in range(r1, r2 + 1):
                                guard.check()
                                for mc in range(c1, c2 + 1):
                                    if (mr, mc) in merged_positions:
                                        raise InputError('MERGED_REGIONS_OVERLAP')
                                    merged_positions.add((mr, mc))
                                    if len(merged_positions) > guard.limits.max_cells:
                                        raise InputError('RESOURCE_LIMIT')
                        rectangle_cells += height * width
                        if rectangle_cells > guard.limits.max_cells:
                            raise InputError('RESOURCE_LIMIT')
                        raw_cells[name] = cells
    except InputError:
        raise
    except Exception:
        raise InputError('PARSE_FAILED') from None
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_links=False)
        guard.check()
        # Resolve worksheets by workbook relationships, not assumed ZIP file order.
        relations = ET.fromstring(members['xl/_rels/workbook.xml.rels'])
        targets = {e.get('Id'): e.get('Target') for e in relations if e.get('TargetMode') != 'External'}
        sheet_elements = list(ET.fromstring(members['xl/workbook.xml']).iter(_XML_NS + 'sheet'))
        xml_paths = []
        for e in sheet_elements:
            target = targets[e.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
            path = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            if path not in raw_cells:
                raise InputError('UNSUPPORTED_FORMAT')
            xml_paths.append(path)
        if len(workbook.worksheets) != len(xml_paths):
            raise InputError('UNSUPPORTED_FORMAT')
        result = []
        total = 0
        for index, sheet in enumerate(workbook.worksheets):
            guard.check()
            raw = raw_cells[xml_paths[index]]
            height = max((r for r, _ in raw), default=0)
            width = max((c for _, c in raw), default=0)
            for merged in sheet.merged_cells.ranges:
                height = max(height, merged.max_row)
                width = max(width, merged.max_col)
            total += height * width
            if height > guard.limits.max_rows or width > guard.limits.max_columns or total > guard.limits.max_cells:
                raise InputError('RESOURCE_LIMIT')
            rows = []
            for row_num in range(1, height + 1):
                row = []
                for col in range(1, width + 1):
                    cell = sheet.cell(row_num, col)
                    raw_type, raw_value, is_formula = raw.get((row_num, col), ('missing', None, False))
                    cache = 'not_applicable'
                    channel = 'value'
                    if is_formula:
                        if not isinstance(cell.value, str):
                            raise InputError('UNSUPPORTED_FORMULA')
                        value = cell.value
                        kind = 'formula'
                        channel = 'formula'
                        cache = 'present_unverified' if raw_value is not None else 'missing'
                    elif raw_type == 'missing':
                        value, kind = '', 'missing'
                    elif raw_type in ('s', 'inlineStr', 'str'):
                        value, kind = cell.value or '', 'text'
                    elif raw_type == 'b':
                        if raw_value not in ('0', '1'):
                            raise InputError('PARSE_FAILED')
                        value, kind = ('true' if raw_value == '1' else 'false'), 'boolean'
                    elif raw_type == 'e':
                        raise InputError('CELL_ERROR')  # Do not report error-valued cells as complete.
                    elif raw_type in ('n', 'd'):
                        value, kind = raw_value or '', 'number' if raw_type == 'n' else 'date'
                    else:
                        raise InputError('UNSUPPORTED_CELL_TYPE')
                    guard.text(value)
                    row.append((value, kind, channel, cache, cell.number_format))
                rows.append(tuple(row))
            hidden_columns = []
            for dim in sheet.column_dimensions.values():
                if dim.hidden:
                    hidden_columns.append((dim.min, dim.max))
            info = dict(name=sheet.title, hidden=sheet.sheet_state != 'visible',
                        hidden_rows=[r for r, d in sheet.row_dimensions.items() if d.hidden],
                        hidden_columns=hidden_columns,
                        merged=[(m.min_row, m.max_row, m.min_col, m.max_col) for m in sheet.merged_cells.ranges])
            result.append((f'xlsx:{index}', rows, info))
        workbook.close()
        if not any(rows for _, rows, _ in result):
            raise InputError('EMPTY_ASSET')
        return result, ['NON_CELL_CONTENT_EXCLUDED', 'FORMULA_CACHES_NOT_SCANNED']
    except InputError:
        raise
    except Exception:
        raise InputError('PARSE_FAILED') from None


def parse_table(data: bytes, asset: AssetRef, *, config=ParserConfig(), limits=Limits(),
                scope: Scope | None = None, cancelled: Callable[[], bool] = lambda: False) -> ParseResult:
    guard = Guard(limits, cancelled)
    guard.check()
    verify_asset(data, asset, limits)
    if not data:
        raise InputError('EMPTY_ASSET')
    if asset.format == 'csv':
        tables, warnings = _csv(data, config, guard)
    elif asset.format == 'xlsx':
        tables, warnings = _xlsx(data, config, guard)
    else:
        raise InputError('UNSUPPORTED_FORMAT')
    dimensions = {sid: (len(rows), len(rows[0]) if rows else 0) for sid, rows, _ in tables}
    if scope is None:
        selections = tuple(Selection(sid, ((1, height),), ((1, width),)) for sid, (height, width) in dimensions.items() if height and width)
        scope = Scope(asset.version_id, selections, config.include_header)
    scope = normalize_scope(scope)
    if scope.asset_version_id != asset.version_id:
        raise InputError('ASSET_VERSION_MISMATCH')
    if scope.include_header != config.include_header:
        raise InputError('SCOPE_INVALID')
    for selection in scope.selections:
        if selection.sheet_id not in dimensions:
            raise InputError('SCOPE_INVALID')
        height, width = dimensions[selection.sheet_id]
        if selection.rows[-1][1] > height or selection.columns[-1][1] > width:
            raise InputError('SCOPE_INVALID')
    selected = {s.sheet_id: s for s in scope.selections}
    cells, skipped, structures = [], [], []
    target = 0
    for sid, rows, info in tables:
        guard.check()
        structures.append(dict(sheet_id=sid, rows=len(rows), columns=len(rows[0]) if rows else 0, **info))
        if sid not in selected:
            continue
        selection = selected[sid]
        aliases = {}
        for r1, r2, c1, c2 in info.get('merged', []):
            for mr in range(r1, r2 + 1):
                guard.check()
                for mc in range(c1, c2 + 1):
                    if (mr, mc) != (r1, c1):
                        aliases[mr, mc] = (r1, c1)
        if config.header_row is not None and config.header_row > len(rows):
            raise InputError('HEADER_OUT_OF_RANGE')
        for r, row in enumerate(rows, 1):
            if not _within(r, selection.rows):
                continue
            for c, value in enumerate(row, 1):
                if not _within(c, selection.columns):
                    continue
                guard.check()
                reason = None
                if r == config.header_row and not config.include_header:
                    reason = 'HEADER_EXCLUDED'
                if config.hidden_policy == 'exclude' and (info.get('hidden') or r in info.get('hidden_rows', []) or any(a <= c <= b for a, b in info.get('hidden_columns', []))):
                    reason = 'HIDDEN_EXCLUDED'
                if (r, c) in aliases:
                    r1, c1 = aliases[r, c]
                    if not (_within(r1, selection.rows) and _within(c1, selection.columns)):
                        raise InputError('SCOPE_MERGE_ANCHOR_MISSING')
                    reason = 'MERGED_ALIAS'
                if reason:
                    skipped.append(dict(sheet_id=sid, row=r, column=c, reason=reason))
                    continue
                target += 1
                if asset.format == 'csv':
                    cells.append(Cell(sid, r, c, value, 'text', normalization=config.normalization))
                else:
                    text, kind, channel, cache, number_format = value
                    cells.append(Cell(sid, r, c, text, kind, channel, cache, number_format, config.normalization))
                    if cache == 'missing':
                        warnings.append('FORMULA_CACHE_MISSING')
                if config.normalization != 'identity-v1':
                    guard.text(cells[-1].normalized_text)
                    if not cells[-1].location_reliable:
                        warnings.append('CHARACTER_MAPPING_UNRELIABLE')
    # Exclusions are outside the normalized eligible denominator, and individually enumerable.
    parse_coverage = dict(status='complete', unit='cell', target_count=target, processed_count=len(cells),
                          failed_count=0, skipped_count=0, ratio=len(cells)/target if target else None,
                          excluded_locations=skipped, processed_locations=[(x.sheet_id, x.row, x.column) for x in cells])
    not_run = dict(status='not_run', target_count=None, processed_count=0, ratio=None, reason='DETECTION_NOT_STARTED')
    coverage = dict(parse=parse_coverage, rule=not_run, semantic=not_run, combination=not_run)
    cfg = dict(parser=asdict(config), limits=asdict(limits), unicode_version=unicodedata.unidata_version)
    return ParseResult(asset.version_id, asset.content_hash, PARSER_VERSION,
                       json.dumps(cfg, ensure_ascii=False, sort_keys=True), fingerprint(cfg), scope,
                       fingerprint(asdict(scope)), tuple(cells), json.dumps(structures, ensure_ascii=False),
                       json.dumps(coverage, ensure_ascii=False), tuple(sorted(set(warnings))),
                       ('comments', 'hyperlink_targets', 'textboxes', 'images', 'formula_cached_values') if asset.format == 'xlsx' else ())
