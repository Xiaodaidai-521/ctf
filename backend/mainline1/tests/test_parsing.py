import hashlib
import io
import json
import unittest
import zipfile
from dataclasses import replace
from unittest.mock import patch

from openpyxl import Workbook

from mainline1.parsing import (
    AssetRef, Cell, InputError, Limits, ParserConfig, Scope, Selection,
    fingerprint, normalize_scope, parse_table, read_private,
)


def asset(data, format='csv'):
    return AssetRef('version-one', 'synthetic-domain', 'sha256:' + hashlib.sha256(data).hexdigest(), format, len(data), 'private:synthetic')


def workbook_bytes(configure):
    book = Workbook()
    configure(book)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def rewrite_zip(data, name, transform):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as old, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as new:
        for info in old.infolist():
            value = old.read(info)
            new.writestr(info.filename, transform(value) if info.filename == name else value)
    return output.getvalue()


class ParsingTests(unittest.TestCase):
    def test_nfkc_maps_fullwidth_emoji_and_combining_to_source(self):
        text = 'Ａ😀e\u0301ﬃ'
        data = ('header\n' + text + '\n').encode('utf-8')
        result = parse_table(data, asset(data), config=ParserConfig(normalization='nfkc-mapped-v1'))
        cell = result.cells[0]
        self.assertEqual(cell.normalized_text, 'A😀éffi')
        self.assertEqual(cell.source_span(0, 1), (0, 1))
        self.assertEqual(cell.source_span(1, 2), (1, 2))
        self.assertEqual(cell.source_span(2, 3), (2, 4))
        self.assertEqual(cell.source_span(3, 6), (4, 5))
        self.assertCode('LOCATION_MISMATCH', cell.source_span, 3, 4)
        self.assertEqual(cell.source_text, text)

    def test_unproven_cross_cluster_mapping_disables_character_operations(self):
        cell = Cell('csv:0', 2, 1, '\u1100\u1161', 'text', normalization='nfkc-mapped-v1')
        self.assertEqual(cell.normalized_text, '가')
        self.assertFalse(cell.location_reliable)
        self.assertCode('LOCATION_MISMATCH', cell.source_span, 0, 1)

    def test_normalization_expansion_obeys_cell_budget(self):
        data = 'h\nﬃ\n'.encode('utf-8')
        self.assertCode('RESOURCE_LIMIT', parse_table, data, asset(data),
                        config=ParserConfig(normalization='nfkc-mapped-v1'),
                        limits=Limits(max_cell_codepoints=2))

    def assertCode(self, code, function, *args, **kwargs):
        with self.assertRaises(InputError) as found:
            function(*args, **kwargs)
        self.assertEqual(found.exception.code, code)
        self.assertEqual(str(found.exception), code)

    def parse(self, data, **kwargs):
        return parse_table(data, asset(data), **kwargs)

    def test_csv_multiline_quotes_duplicate_empty_headers_and_logical_records(self):
        data = '重复,重复,\r\n"甲\r\n乙","a""b",\r\n'.encode()
        result = self.parse(data)
        self.assertEqual([(c.row, c.column, c.source_text) for c in result.cells], [(2, 1, '甲\r\n乙'), (2, 2, 'a"b'), (2, 3, '')])
        with_header = self.parse(data, config=ParserConfig(include_header=True))
        self.assertEqual([c.source_text for c in with_header.cells[:3]], ['重复', '重复', ''])
        self.assertEqual(result.scope.selections[0].rows, ((1, 2),))

    def test_csv_explicit_encoding_dialect_and_no_header(self):
        data = '中文;编号\r内容;0012'.encode('gb18030')
        result = self.parse(data, config=ParserConfig(encoding='gb18030', delimiter=';', header_row=None))
        self.assertEqual([c.source_text for c in result.cells], ['中文', '编号', '内容', '0012'])
        self.assertCode('ENCODING_INVALID', self.parse, data)

    def test_bom_requires_bom_aware_encoding(self):
        data = b'\xef\xbb\xbfh\nx\n'
        self.assertEqual(self.parse(data).cells[0].source_text, 'x')
        self.assertCode('ENCODING_INVALID', self.parse, data, config=ParserConfig(encoding='utf-8'))

    def test_empty_and_malformed_csv_are_not_success(self):
        for raw, code in [(b'', 'EMPTY_ASSET'), (b'h\n"unterminated', 'PARSE_FAILED'), (b'a,b\nx\n', 'PARSE_FAILED'), (b'\n', 'PARSE_FAILED'), (b'a\x00b', 'PARSE_FAILED')]:
            with self.subTest(raw=raw):
                self.assertCode(code, self.parse, raw)

    def test_limits_file_rows_columns_cells_text(self):
        data = b'a,b\nx,y\n'
        for limits in [Limits(max_bytes=2), Limits(max_rows=1), Limits(max_columns=1), Limits(max_cells=3)]:
            self.assertCode('RESOURCE_LIMIT', self.parse, data, limits=limits)
        self.assertCode('RESOURCE_LIMIT', self.parse, b'h\nlong\n', limits=Limits(max_cell_codepoints=2))

    def test_invalid_configuration_does_not_guess(self):
        for kwargs in [dict(encoding='latin-1'), dict(delimiter=',' ,quotechar=','), dict(formula_policy='evaluate'), dict(normalization='NFKC'), dict(header_row=0)]:
            with self.subTest(kwargs=kwargs), self.assertRaises(InputError):
                ParserConfig(**kwargs)

    def test_identity_mapping_emoji_combining_fullwidth_and_space(self):
        text = ' A😀e\u0301Ａ\nZ '
        result = self.parse(('h\n"'+text+'"\n').encode())
        cell = result.cells[0]
        self.assertEqual(cell.normalized_text, text)
        self.assertEqual(cell.source_span(2, 3), (2, 3))
        self.assertEqual(cell.source_text[2:3], '😀')
        self.assertEqual(cell.source_text[3:5], 'e\u0301')
        self.assertCode('LOCATION_MISMATCH', cell.source_span, 0, len(text)+1)
        self.assertNotIn(text, repr(cell))
        self.assertNotIn(text, repr(result))

    def test_normalized_scope_hash_and_version_binding(self):
        a = Scope('version-one', (Selection('csv:0', ((2, 2), (1, 1)), ((2, 2), (1, 1))),))
        b = Scope('version-one', (Selection('csv:0', ((1, 2),), ((1, 2),)),))
        self.assertEqual(normalize_scope(a), b)
        self.assertEqual(self.parse(b'h,h\nx,y\n', scope=a).scope_hash, self.parse(b'h,h\nx,y\n', scope=b).scope_hash)
        self.assertCode('ASSET_VERSION_MISMATCH', self.parse, b'h,h\nx,y\n', scope=replace(a, asset_version_id='old'))
        self.assertCode('SCOPE_INVALID', self.parse, b'h\nx\n', scope=Scope('version-one', (Selection('csv:0', ((1, 3),), ((1, 1),)),)))

    def test_csv_scoped_projection_keeps_original_coordinates(self):
        selected = Scope('version-one', (Selection('csv:0', ((3, 3),), ((2, 2),)),))
        result = self.parse(b'a,b\n1,2\n3,4\n', scope=selected)
        self.assertEqual([(c.row, c.column, c.source_text) for c in result.cells], [(3, 2, '4')])

    def test_config_snapshot_is_immutable_and_covers_parser_budget(self):
        result = self.parse(b'h\nx\n')
        self.assertEqual(json.loads(result.config_json)['parser']['normalization'], 'identity-v1')
        other = self.parse(b'h\nx\n', limits=Limits(max_rows=20))
        self.assertNotEqual(result.config_hash, other.config_hash)
        with self.assertRaises(Exception):
            result.cells[0].row = 999
        self.assertCode('FINGERPRINT_TYPE_UNSUPPORTED', fingerprint, {'float': 0.1})

    def test_not_run_detection_coverage_and_explicit_exclusions(self):
        result = self.parse(b'a,b\nx,y\n')
        cov = json.loads(result.coverage_json)
        self.assertEqual(cov['parse']['processed_count'], 2)
        self.assertEqual(len(cov['parse']['excluded_locations']), 2)
        for layer in ('rule', 'semantic', 'combination'):
            self.assertEqual(cov[layer]['status'], 'not_run')
            self.assertIsNone(cov[layer]['ratio'])

    def test_cancel_before_and_during_parse(self):
        self.assertCode('TASK_CANCELLED', self.parse, b'h\nx\n', cancelled=lambda: True)
        calls = []
        def cancel():
            calls.append(1)
            return len(calls) > 3
        self.assertCode('TASK_CANCELLED', self.parse, b'h\nx\ny\nz\n', cancelled=cancel)

    def test_cooperative_time_budget(self):
        with patch('mainline1.parsing.time.monotonic', side_effect=[0, 301]):
            self.assertCode('RESOURCE_LIMIT', self.parse, b'h\nx\n')

    def test_hash_mismatch_format_and_deleted_state(self):
        data = b'h\nx\n'; ref = asset(data)
        self.assertCode('ASSET_VERSION_MISMATCH', parse_table, data+b'y', ref)
        self.assertCode('ASSET_UNAVAILABLE', parse_table, data, replace(ref, state='deleted'))
        self.assertCode('UNSUPPORTED_FORMAT', parse_table, data, replace(ref, format='xlsm'))

    def test_private_source_must_reauthorize_and_stream_is_closed(self):
        data = b'h\nx\n'; ref = asset(data)
        class Source:
            def open_version(self, ref, *, principal, domain_id):
                self.args = (principal, domain_id)
                self.stream = io.BytesIO(data)
                return self.stream
        source = Source()
        self.assertEqual(read_private(source, ref, principal='p', domain_id=ref.domain_id, limits=Limits()), data)
        self.assertTrue(source.stream.closed)
        self.assertEqual(source.args, ('p', 'synthetic-domain'))
        self.assertCode('OBJECT_NOT_FOUND', read_private, source, ref, principal='p', domain_id='other', limits=Limits())
        self.assertCode('RESOURCE_LIMIT', read_private, source, replace(ref, byte_size=1), principal='p', domain_id=ref.domain_id, limits=Limits(max_bytes=2))

    def test_xlsx_sheets_hidden_merge_and_formula_cache_missing(self):
        def configure(w):
            s=w.active;s.title='Duplicate labels';s.append(['h','h']);s.append(['😀','=1+1']);s.merge_cells('A3:B3');s['A3']='anchor';s.row_dimensions[2].hidden=True;s.column_dimensions['B'].hidden=True
            other=w.create_sheet('hidden');other.sheet_state='hidden';other.append(['h']);other.append(['secret-synthetic'])
        data=workbook_bytes(configure)
        result=parse_table(data,asset(data,'xlsx'))
        formula=next(c for c in result.cells if c.channel=='formula')
        self.assertEqual((formula.source_text,formula.cache_status),('=1+1','missing'))
        self.assertIn('FORMULA_CACHE_MISSING',result.warnings)
        self.assertIn('hyperlink_targets',result.excluded_contents)
        self.assertEqual(sum(c.source_text=='anchor' for c in result.cells),1)
        self.assertTrue(any(c.sheet_id=='xlsx:1' for c in result.cells))
        excluded=parse_table(data,asset(data,'xlsx'),config=ParserConfig(hidden_policy='exclude'))
        self.assertFalse(any(c.row==2 or c.sheet_id=='xlsx:1' for c in excluded.cells))
        scope=Scope('version-one',(Selection('xlsx:0',((3,3),),((2,2),)),))
        self.assertCode('SCOPE_MERGE_ANCHOR_MISSING',parse_table,data,asset(data,'xlsx'),scope=scope)

    def test_xlsx_cached_value_observed_not_executed_or_trusted(self):
        data=workbook_bytes(lambda w: (w.active.append(['h']),w.active.append(['=2+3'])))
        data=rewrite_zip(data,'xl/worksheets/sheet1.xml',lambda b:b.replace(b'<v></v>',b'<v>999</v>'))
        result=parse_table(data,asset(data,'xlsx'))
        self.assertEqual(result.cells[0].source_text,'=2+3')
        self.assertEqual(result.cells[0].cache_status,'present_unverified')

    def test_xlsx_raw_numeric_boolean_and_empty_types(self):
        def configure(w):
            w.active.append(['a','b','c','d']);w.active.append([12,True,'',None]);w.active['A2'].number_format='0000'
        data=workbook_bytes(configure);result=parse_table(data,asset(data,'xlsx'))
        self.assertEqual([c.value_type for c in result.cells],['number','boolean','text','missing'])
        self.assertEqual([c.source_text for c in result.cells],['12','true','',''])
        self.assertEqual(result.cells[0].number_format,'0000')

    def test_xlsx_budget_rejects_before_openpyxl(self):
        data=workbook_bytes(lambda w: w.active.append(['x']))
        for limits in (Limits(max_zip_entries=1),Limits(max_uncompressed_bytes=10)):
            with patch('openpyxl.load_workbook') as loader:
                self.assertCode('RESOURCE_LIMIT',parse_table,data,asset(data,'xlsx'),limits=limits)
                loader.assert_not_called()
        modified=rewrite_zip(data,'xl/worksheets/sheet1.xml',lambda b:b.replace(b'r="A1"',b'r="Z1"'))
        with patch('openpyxl.load_workbook') as loader:
            self.assertCode('RESOURCE_LIMIT',parse_table,modified,asset(modified,'xlsx'),limits=Limits(max_columns=2))
            loader.assert_not_called()

    def test_xlsx_unsafe_xml_and_archive_paths(self):
        data=workbook_bytes(lambda w: w.active.append(['x']))
        modified=rewrite_zip(data,'xl/worksheets/sheet1.xml',lambda b:b'<!DOCTYPE bad [<!ENTITY x "x">]>'+b)
        self.assertCode('UNSAFE_XML',parse_table,modified,asset(modified,'xlsx'))
        stream=io.BytesIO()
        with zipfile.ZipFile(stream,'w') as archive:archive.writestr('../outside','synthetic')
        self.assertCode('UNSAFE_ARCHIVE',parse_table,stream.getvalue(),asset(stream.getvalue(),'xlsx'))

    def test_empty_xlsx_is_not_success(self):
        data=workbook_bytes(lambda w:None)
        self.assertCode('EMPTY_ASSET',parse_table,data,asset(data,'xlsx'))


if __name__ == '__main__':
    unittest.main()
