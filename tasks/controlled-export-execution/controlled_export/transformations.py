import csv
import io
import zipfile
from dataclasses import dataclass

from django.conf import settings

from .authorization import digest_json, sha256_bytes
from .errors import ContractError, UnsupportedContent


FORMULA_PREFIXES = ("=", "+", "-", "@")
ALLOWED_OPERATIONS = {"mask_columns", "drop_columns", "sanitize_formulas"}
REJECTED_XLSX_MEMBERS = ("xl/externalLinks/", "xl/embeddings/", "xl/vbaProject", "xl/connections")


@dataclass(frozen=True)
class TransformResult:
    content: bytes
    output_format: str
    suffix: str
    processed_records: int
    output_records: int
    output_hash: str
    scope_hash: str
    plan_hash: str


def neutralize_formula(value):
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _validate_columns(columns, field):
    if not isinstance(columns, list) or not columns:
        raise ContractError(f"{field} must be a non-empty list", code="PLAN_SCOPE_INVALID")
    if any(not isinstance(item, str) or not item or len(item) > 200 for item in columns):
        raise ContractError(f"{field} contains an invalid column", code="PLAN_SCOPE_INVALID")
    if len(set(columns)) != len(columns):
        raise ContractError(f"{field} contains duplicate columns", code="PLAN_SCOPE_INVALID")
    return columns


def validate_plan(plan, input_format, effective_action):
    if not isinstance(plan, dict) or plan.get("version") != "1":
        raise ContractError("A version 1 transformation plan is required", code="PLAN_INVALID")
    if plan.get("format") != input_format:
        raise ContractError("Plan format does not match input format", code="PLAN_FORMAT_MISMATCH")
    scope = plan.get("scope")
    if not isinstance(scope, dict):
        raise ContractError("Plan scope must be an object", code="PLAN_SCOPE_INVALID")
    _validate_columns(scope.get("columns"), "scope.columns")
    if input_format == "xlsx":
        sheets = scope.get("worksheets")
        if not isinstance(sheets, list) or not sheets or any(not isinstance(item, str) or not item for item in sheets):
            raise ContractError("scope.worksheets must be a non-empty list", code="PLAN_SCOPE_INVALID")
        if len(set(sheets)) != len(sheets):
            raise ContractError("scope.worksheets contains duplicates", code="PLAN_SCOPE_INVALID")
    operations = plan.get("operations", [])
    if not isinstance(operations, list):
        raise ContractError("Plan operations must be a list", code="PLAN_INVALID")
    has_protection = False
    for operation in operations:
        if not isinstance(operation, dict) or operation.get("type") not in ALLOWED_OPERATIONS:
            raise UnsupportedContent("Plan contains an unsupported operation", code="PLAN_OPERATION_UNSUPPORTED")
        if operation["type"] in {"mask_columns", "drop_columns"}:
            _validate_columns(operation.get("columns"), f"{operation['type']}.columns")
            has_protection = True
        if operation["type"] == "mask_columns":
            if operation.get("strategy", "fixed") != "fixed":
                raise UnsupportedContent("Only fixed-value masking is approved", code="MASK_STRATEGY_UNSUPPORTED")
            replacement = operation.get("value", "***")
            if not isinstance(replacement, str) or len(replacement) > 256:
                raise ContractError("Mask replacement is invalid", code="MASK_VALUE_INVALID")
    if effective_action == "mask" and not has_protection:
        raise ContractError("Mask authorization requires a masking or removal operation", code="MASK_PLAN_REQUIRED")
    return scope, operations


def _operation_map(scope_columns, operations):
    drops = set()
    masks = {}
    for operation in operations:
        columns = operation.get("columns", [])
        unknown = set(columns).difference(scope_columns)
        if unknown:
            raise ContractError(f"Operation columns outside scope: {sorted(unknown)}", code="PLAN_SCOPE_MISMATCH")
        if operation["type"] == "drop_columns":
            drops.update(columns)
        elif operation["type"] == "mask_columns":
            replacement = operation.get("value", "***")
            for column in columns:
                if column in masks and masks[column] != replacement:
                    raise ContractError("Conflicting mask operations", code="PLAN_OPERATION_CONFLICT")
                masks[column] = replacement
    if drops.intersection(masks):
        raise ContractError("A column cannot be both masked and dropped", code="PLAN_OPERATION_CONFLICT")
    output_columns = [column for column in scope_columns if column not in drops]
    if not output_columns:
        raise ContractError("Plan would remove every column", code="PLAN_EMPTY_OUTPUT")
    return output_columns, masks


def transform_csv(content, plan, effective_action):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UnsupportedContent("CSV input must be UTF-8", code="CSV_ENCODING_UNSUPPORTED") from exc
    scope, operations = validate_plan(plan, "csv", effective_action)
    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        header = next(reader)
    except StopIteration as exc:
        raise UnsupportedContent("CSV input is empty", code="CSV_EMPTY") from exc
    if not header or len(header) > settings.CONTROLLED_EXPORT_MAX_COLUMNS:
        raise UnsupportedContent("CSV column limit exceeded", code="CSV_COLUMN_LIMIT")
    if any(not name for name in header) or len(set(header)) != len(header):
        raise UnsupportedContent("CSV headers must be unique and non-empty", code="CSV_HEADER_INVALID")
    scope_columns = scope["columns"]
    missing = set(scope_columns).difference(header)
    if missing:
        raise ContractError(f"Authorized columns are missing: {sorted(missing)}", code="PLAN_SCOPE_MISMATCH")
    output_columns, masks = _operation_map(scope_columns, operations)
    indexes = {name: header.index(name) for name in scope_columns}
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(output_columns)
    row_count = 0
    for row in reader:
        row_count += 1
        if row_count > settings.CONTROLLED_EXPORT_MAX_ROWS:
            raise UnsupportedContent("CSV row limit exceeded", code="CSV_ROW_LIMIT")
        if len(row) != len(header):
            raise UnsupportedContent("CSV rows must match the header width", code="CSV_WIDTH_INVALID")
        writer.writerow(
            neutralize_formula(masks.get(column, row[indexes[column]]))
            for column in output_columns
        )
    result = output.getvalue().encode("utf-8")
    return TransformResult(
        content=result,
        output_format="csv",
        suffix=".csv",
        processed_records=row_count,
        output_records=row_count,
        output_hash=sha256_bytes(result),
        scope_hash=digest_json(scope),
        plan_hash=digest_json(plan),
    )


def _preflight_xlsx(content):
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise UnsupportedContent("XLSX input is not a valid ZIP package", code="XLSX_INVALID") from exc
    with archive:
        members = archive.infolist()
        if len(members) > 5000:
            raise UnsupportedContent("XLSX package has too many members", code="XLSX_RESOURCE_LIMIT")
        expanded = sum(member.file_size for member in members)
        if expanded > settings.CONTROLLED_EXPORT_MAX_INPUT_BYTES * 20:
            raise UnsupportedContent("XLSX expanded size limit exceeded", code="XLSX_RESOURCE_LIMIT")
        for member in members:
            normalized = member.filename.replace("\\", "/")
            if normalized.startswith("/") or ".." in normalized.split("/"):
                raise UnsupportedContent("XLSX package contains an unsafe path", code="XLSX_PATH_INVALID")
            if any(normalized.startswith(prefix) for prefix in REJECTED_XLSX_MEMBERS):
                raise UnsupportedContent("XLSX contains unsupported active or external content", code="XLSX_ACTIVE_CONTENT")


def transform_xlsx(content, plan, effective_action):
    _preflight_xlsx(content)
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise UnsupportedContent("openpyxl is required for XLSX processing", code="XLSX_PROVIDER_UNAVAILABLE") from exc

    scope, operations = validate_plan(plan, "xlsx", effective_action)
    source = load_workbook(io.BytesIO(content), read_only=False, data_only=False, keep_links=False)
    selected_sheets = scope["worksheets"]
    missing_sheets = set(selected_sheets).difference(source.sheetnames)
    if missing_sheets:
        raise ContractError(f"Authorized worksheets are missing: {sorted(missing_sheets)}", code="PLAN_SCOPE_MISMATCH")
    output_columns, masks = _operation_map(scope["columns"], operations)
    destination = Workbook(write_only=False)
    destination.remove(destination.active)
    total_rows = 0
    try:
        for sheet_name in selected_sheets:
            sheet = source[sheet_name]
            if sheet.sheet_state != "visible":
                raise UnsupportedContent("Selected worksheet is hidden", code="XLSX_HIDDEN_CONTENT")
            if sheet.merged_cells.ranges:
                raise UnsupportedContent("Merged cells are not supported", code="XLSX_MERGED_CONTENT")
            if any(dimension.hidden for dimension in sheet.row_dimensions.values()):
                raise UnsupportedContent("Selected worksheet contains hidden rows", code="XLSX_HIDDEN_CONTENT")
            if any(dimension.hidden for dimension in sheet.column_dimensions.values()):
                raise UnsupportedContent("Selected worksheet contains hidden columns", code="XLSX_HIDDEN_CONTENT")
            rows = sheet.iter_rows()
            try:
                header_cells = next(rows)
            except StopIteration as exc:
                raise UnsupportedContent("Selected worksheet is empty", code="XLSX_EMPTY") from exc
            header = [cell.value for cell in header_cells]
            if any(not isinstance(value, str) or not value for value in header) or len(set(header)) != len(header):
                raise UnsupportedContent("XLSX headers must be unique non-empty strings", code="XLSX_HEADER_INVALID")
            missing = set(scope["columns"]).difference(header)
            if missing:
                raise ContractError(f"Authorized columns are missing: {sorted(missing)}", code="PLAN_SCOPE_MISMATCH")
            indexes = {name: header.index(name) for name in scope["columns"]}
            target = destination.create_sheet(sheet_name)
            target.append(output_columns)
            for source_row in rows:
                total_rows += 1
                if total_rows > settings.CONTROLLED_EXPORT_MAX_ROWS:
                    raise UnsupportedContent("XLSX row limit exceeded", code="XLSX_ROW_LIMIT")
                for cell in source_row:
                    if cell.data_type == "f" or cell.comment is not None or cell.hyperlink is not None:
                        raise UnsupportedContent(
                            "Selected worksheet contains formulas, comments, or hyperlinks",
                            code="XLSX_ACTIVE_CONTENT",
                        )
                target.append(
                    [
                        neutralize_formula(masks.get(column, source_row[indexes[column]].value))
                        for column in output_columns
                    ]
                )
    finally:
        source.close()
    output = io.BytesIO()
    destination.save(output)
    result = output.getvalue()
    _preflight_xlsx(result)
    return TransformResult(
        content=result,
        output_format="xlsx",
        suffix=".xlsx",
        processed_records=total_rows,
        output_records=total_rows,
        output_hash=sha256_bytes(result),
        scope_hash=digest_json(scope),
        plan_hash=digest_json(plan),
    )


def transform(content, filename, plan, effective_action):
    if not isinstance(content, bytes) or not content:
        raise UnsupportedContent("Input content is empty", code="INPUT_EMPTY")
    if len(content) > settings.CONTROLLED_EXPORT_MAX_INPUT_BYTES:
        raise UnsupportedContent("Input size limit exceeded", code="INPUT_SIZE_LIMIT")
    lowered = filename.lower()
    if lowered.endswith(".csv"):
        return transform_csv(content, plan, effective_action)
    if lowered.endswith(".xlsx"):
        return transform_xlsx(content, plan, effective_action)
    raise UnsupportedContent("Only CSV and XLSX are supported", code="FORMAT_UNSUPPORTED")
