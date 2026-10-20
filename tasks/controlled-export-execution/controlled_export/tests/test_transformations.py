import csv
import io

from django.test import SimpleTestCase, override_settings
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from controlled_export.errors import ContractError, UnsupportedContent
from controlled_export.transformations import transform_csv, transform_xlsx

from .helpers import csv_bytes, csv_plan


@override_settings(CONTROLLED_EXPORT_MAX_INPUT_BYTES=1024 * 1024, CONTROLLED_EXPORT_MAX_ROWS=100)
class TransformationTests(SimpleTestCase):
    def test_csv_is_scoped_masked_and_formula_neutralized(self):
        result = transform_csv(csv_bytes(), csv_plan(), "mask")
        self.assertEqual(
            result.content.decode("utf-8").splitlines(),
            ["name,email,note", "Alice,***,'=1+1"],
        )
        self.assertEqual(result.output_records, 1)

    def test_csv_rejects_columns_outside_authorized_scope(self):
        plan = csv_plan()
        plan["operations"][0]["columns"] = ["secret"]
        with self.assertRaises(ContractError):
            transform_csv(csv_bytes(), plan, "mask")

    def _workbook_bytes(self, *, hidden=False, formula=False, hyperlink=False):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Export"
        sheet.append(["name", "email", "note"])
        sheet.append(["Synthetic User", "user@example.invalid", "=2+2" if formula else "safe"])
        if hidden:
            sheet.column_dimensions["B"].hidden = True
        if hyperlink:
            sheet["C2"].hyperlink = "https://example.invalid"
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _xlsx_plan(self):
        return {
            "version": "1",
            "format": "xlsx",
            "scope": {"worksheets": ["Export"], "columns": ["name", "email", "note"]},
            "operations": [
                {"type": "mask_columns", "columns": ["email"], "strategy": "fixed", "value": "***"},
                {"type": "sanitize_formulas"},
            ],
        }

    def test_xlsx_is_rebuilt_with_only_selected_content(self):
        result = transform_xlsx(self._workbook_bytes(), self._xlsx_plan(), "mask")
        workbook = load_workbook(io.BytesIO(result.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Export"])
        self.assertEqual(workbook["Export"]["B2"].value, "***")
        self.assertEqual(workbook["Export"]["C2"].value, "safe")
        workbook.close()

    def test_xlsx_rejects_hidden_formula_and_hyperlink_channels(self):
        for options in ({"hidden": True}, {"formula": True}, {"hyperlink": True}):
            with self.subTest(options=options), self.assertRaises(UnsupportedContent):
                transform_xlsx(self._workbook_bytes(**options), self._xlsx_plan(), "mask")


@override_settings(CONTROLLED_EXPORT_MAX_INPUT_BYTES=1024 * 1024, CONTROLLED_EXPORT_MAX_ROWS=100)
class HeaderFormulaTests(SimpleTestCase):
    def _csv(self, rows):
        output = io.StringIO(newline="")
        csv.writer(output, lineterminator="\n").writerows(rows)
        return output.getvalue().encode("utf-8")

    def _csv_plan(self, columns):
        return {
            "version": "1",
            "format": "csv",
            "scope": {"columns": list(columns)},
            "operations": [{"type": "sanitize_formulas"}],
        }

    def test_csv_formula_like_headers_are_neutralized(self):
        for header in ("=1+1", "+1+1", "-1+1", "@SUM(A1)"):
            with self.subTest(header=header):
                result = transform_csv(
                    self._csv([[header, "note"], ["value", "safe"]]),
                    self._csv_plan([header, "note"]),
                    "allow",
                )
                rows = list(csv.reader(io.StringIO(result.content.decode("utf-8"))))
                self.assertEqual(rows[0], ["'" + header, "note"])
                self.assertEqual(rows[1], ["value", "safe"])

    def test_csv_header_only_file_still_neutralizes(self):
        result = transform_csv(self._csv([["=1+1"]]), self._csv_plan(["=1+1"]), "allow")
        self.assertEqual(result.content.decode("utf-8").splitlines(), ["'=1+1"])
        self.assertEqual(result.output_records, 0)

    def test_csv_rejects_headers_that_collide_after_neutralization(self):
        with self.assertRaises(UnsupportedContent):
            transform_csv(
                self._csv([["=x", "'=x"], ["a", "b"]]),
                self._csv_plan(["=x", "'=x"]),
                "allow",
            )

    def _xlsx(self, header, *, comment_column=None, hyperlink_column=None, formula_column=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Export"
        sheet.append(header)
        sheet.append(["value"] * len(header))
        if comment_column is not None:
            sheet.cell(row=1, column=comment_column + 1).comment = Comment("note", "author")
        if hyperlink_column is not None:
            sheet.cell(row=1, column=hyperlink_column + 1).hyperlink = "https://example.invalid"
        if formula_column is not None:
            sheet.cell(row=1, column=formula_column + 1).value = "=1+1"
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _xlsx_plan(self, columns):
        return {
            "version": "1",
            "format": "xlsx",
            "scope": {"worksheets": ["Export"], "columns": list(columns)},
            "operations": [{"type": "sanitize_formulas"}],
        }

    def test_xlsx_text_formula_header_is_neutralized_not_written_as_formula(self):
        header = ["+1", "note"]
        result = transform_xlsx(self._xlsx(header), self._xlsx_plan(header), "allow")
        workbook = load_workbook(io.BytesIO(result.content), data_only=False)
        cell = workbook["Export"]["A1"]
        self.assertEqual(cell.value, "'+1")
        self.assertNotEqual(cell.data_type, "f")
        workbook.close()

    def test_xlsx_rejects_active_header_cells(self):
        cases = (
            ({"formula_column": 0}, ["=1+1", "note"]),
            ({"comment_column": 0}, ["name", "note"]),
            ({"hyperlink_column": 0}, ["name", "note"]),
        )
        for options, header in cases:
            with self.subTest(options=options), self.assertRaises(UnsupportedContent):
                transform_xlsx(
                    self._xlsx(header, **options),
                    self._xlsx_plan(header),
                    "allow",
                )

    def test_xlsx_rejects_active_header_outside_authorized_scope(self):
        with self.assertRaises(UnsupportedContent):
            transform_xlsx(
                self._xlsx(["name", "secret", "note"], comment_column=1),
                self._xlsx_plan(["name", "note"]),
                "allow",
            )

    def test_xlsx_rejects_headers_that_collide_after_neutralization(self):
        with self.assertRaises(UnsupportedContent):
            transform_xlsx(
                self._xlsx(["=x", "'=x"]),
                self._xlsx_plan(["=x", "'=x"]),
                "allow",
            )
